"""
從公司範本「樣本-出差報告暨旅費核銷.xlsx」抽出樣式規格 → template-spec.js

用法：
    python tools/extract_template_spec.py "樣本-出差報告暨旅費核銷.xlsx"

只抽「樣式」與「固定標籤文字」，不抽範本裡的範例資料（人名、金額、備註）。
公司範本若改版，重跑本腳本即可，app.js 不用動。
"""
import io
import json
import re
import sys
import zipfile
from pathlib import Path

import openpyxl
from openpyxl.styles.colors import COLOR_INDEX
from openpyxl.utils import get_column_letter

MAIN = '2. 出差報告暨旅費核銷明細表'
RATE = '1. 匯率表輸入'
COST = '出差費用種類'          # 整頁照抄（公司規定文字）
DAILY = '出差日報及交際費明細'  # 表頭照抄，資料列由 app.js 依出差日期產生
COLS = 17  # A..Q

# 範本列 → 角色。app.js 依角色動態排版（明細列數、幣別列數可變）
ROLES = {
    'detail': 28,       # 一般明細列
    'deduct': 40,       # 超限扣回列（紅字）
    'subtotal': 41,     # 小計
    'totHead': 42,      # 檢附憑證/旅費總計 表頭
    'totFirst': 43,     # 第一個幣別
    'totMid': 44,       # 中間幣別
    'totLast': 47,      # 最後一個幣別
    'totSum': 48,       # 本幣合計
    'note1': 49,        # 使用備註（出差費用種類說明）
    'supp': 50,         # 補充說明
    'note2': 51,        # 使用備註（淺黃底…）
    'gap': 52,
    'entTitle': 53,     # 交際費標題
    'entDetail': 54,    # 交際費明細列
    'entHead': 58,
    'entFirst': 59,
    'entMid': 60,
    'entLast': 61,
    'entSum': 62,       # 交際費合計
    'grand': 63,        # 出差&交際費總合
}
# 這些角色列只保留固定標籤文字；其他格(範例資料)不帶
LABEL_KEEP = {
    'subtotal': {'L'},
    'totHead': {'A', 'B', 'C', 'D', 'F', 'H', 'J'},
    'totSum': {'D'},
    'note1': {'A', 'B'},
    'supp': {'A'},
    'note2': {'A', 'B'},
    'entTitle': {'A'},
    'entHead': {'A', 'B', 'C', 'D', 'F', 'H', 'J'},
    'entSum': {'D'},
    'grand': {'A'},
}
# 表頭 1~27 列中屬於「資料」的格子（不帶範例值，由 app.js 填）
HEADER_DATA_CELLS = {'A1', 'A3', 'H3', 'B4', 'H4', 'B5', 'C6', 'F6', 'L6', 'A9', 'A16'}

# Office 2013-2022 佈景主題色（範本使用此主題）
THEME = ['FFFFFF', '000000', 'E7E6E6', '44546A', '4472C4', 'ED7D31', 'A5A5A5', 'FFC000', '5B9BD5', '70AD47']


def tint(hex6, t):
    import colorsys
    r, g, b = (int(hex6[i:i + 2], 16) / 255 for i in (0, 2, 4))
    h, l, s = colorsys.rgb_to_hls(r, g, b)
    l = l * (1 + t) if t < 0 else l * (1 - t) + t
    r, g, b = colorsys.hls_to_rgb(h, l, s)
    return ''.join(f'{round(x * 255):02X}' for x in (r, g, b))


def color(c):
    if c is None:
        return None
    if c.type == 'rgb' and isinstance(c.rgb, str):
        return 'FF' + c.rgb[-6:]
    if c.type == 'theme' and isinstance(c.theme, int):
        return 'FF' + tint(THEME[c.theme], c.tint or 0)
    if c.type == 'indexed' and isinstance(c.indexed, int) and c.indexed < len(COLOR_INDEX):
        return 'FF' + COLOR_INDEX[c.indexed][-6:]
    return None


def style_of(c):
    s = {}
    f = c.font
    font = {'name': f.name, 'size': f.sz}
    if f.b:
        font['bold'] = True
    fc = color(f.color)
    if fc and fc != 'FF000000':
        font['color'] = {'argb': fc}
    s['font'] = font
    a = c.alignment
    al = {}
    if a.horizontal:
        al['horizontal'] = a.horizontal
    if a.vertical:
        al['vertical'] = 'middle' if a.vertical == 'center' else a.vertical
    if a.wrap_text:
        al['wrapText'] = True
    if al:
        s['alignment'] = al
    bd = {}
    for side in ('left', 'right', 'top', 'bottom'):
        b = getattr(c.border, side)
        if b is not None and b.style:
            e = {'style': b.style}
            bc = color(b.color)
            if bc and bc != 'FF000000':
                e['color'] = {'argb': bc}
            bd[side] = e
    if bd:
        s['border'] = bd
    if c.fill is not None and c.fill.fill_type == 'solid':
        fg = color(c.fill.fgColor)
        if fg:
            s['fill'] = {'type': 'pattern', 'pattern': 'solid', 'fgColor': {'argb': fg}}
    if c.number_format and c.number_format != 'General':
        s['numFmt'] = c.number_format
    return s


def load_without_merges(path):
    """openpyxl 讀檔時會把合併範圍內(非左上角)格子的樣式丟掉，
    所以另外讀一份「拿掉 <mergeCells>」的版本，取得每一格真正的框線/字型。"""
    src = zipfile.ZipFile(path)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, 'w', zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename.startswith('xl/worksheets/sheet') and item.filename.endswith('.xml'):
                data = re.sub(rb'<mergeCells\b.*?</mergeCells>', b'', data, flags=re.S)
            dst.writestr(item, data)
    buf.seek(0)
    return openpyxl.load_workbook(buf)


def col_widths(ws, n):
    """<col min=1 max=3> 這種群組寬度要展開到每一欄"""
    default = ws.sheet_format.defaultColWidth or 8.875
    out = [default] * n
    for d in ws.column_dimensions.values():
        if d.width and d.customWidth and d.min:
            for i in range(d.min, min(d.max or d.min, n) + 1):
                out[i - 1] = d.width
    return out


def main(path):
    wb_merged = openpyxl.load_workbook(path)   # 只用來讀合併範圍
    wb = load_without_merges(path)            # 讀樣式
    styles, index = [], {}

    def sid(c):
        key = json.dumps(style_of(c), sort_keys=True, ensure_ascii=False)
        if key not in index:
            index[key] = len(styles)
            styles.append(json.loads(key))
        return index[key]

    def row_spec(ws, r, ncols, keep=None, skip=None):
        cells = []
        for i in range(1, ncols + 1):
            c = ws.cell(r, i)
            col = get_column_letter(i)
            v = c.value
            if isinstance(v, str) and v.startswith('='):
                v = None
            if skip and f'{col}{r}' in skip:
                v = None
            if keep is not None and col not in keep and v != '元':
                v = None
            if v is not None and not isinstance(v, (str, int, float)):
                v = None
            cells.append([sid(c), v])
        return {'h': ws.row_dimensions[r].height, 'cells': cells}

    ws = wb[MAIN]
    header = [row_spec(ws, r, COLS, skip=HEADER_DATA_CELLS) for r in range(1, 28)]
    merges = sorted(str(m) for m in wb_merged[MAIN].merged_cells.ranges if m.max_row <= 27)
    roles = {k: row_spec(ws, r, COLS, keep=LABEL_KEEP.get(k, set())) for k, r in ROLES.items()}
    widths = col_widths(ws, COLS)
    ps = ws.page_setup
    m = ws.page_margins
    page = {
        'paperSize': int(ps.paperSize or 9), 'orientation': ps.orientation or 'portrait',
        'scale': int(ps.scale or 100),
        'margins': {'left': m.left, 'right': m.right, 'top': m.top, 'bottom': m.bottom, 'header': m.header, 'footer': m.footer},
    }

    wr = wb[RATE]
    rate = {
        'header': row_spec(wr, 1, 4),
        'row': row_spec(wr, 2, 4, keep=set()),
        'widths': col_widths(wr, 4),
    }

    # 出差費用種類：整頁照抄
    wc = wb[COST]
    cost = {
        'rows': [row_spec(wc, r, 3) for r in range(1, wc.max_row + 1)],
        'merges': sorted(str(x) for x in wb_merged[COST].merged_cells.ranges),
        'widths': col_widths(wc, 3),
    }

    # 出差日報及交際費明細：標題列、表頭列、資料列樣式、交際費總計樣式
    wd = wb[DAILY]
    dps = wd.page_setup
    daily = {
        'title': row_spec(wd, 1, 13),
        'head': row_spec(wd, 2, 13),
        'row': row_spec(wd, 3, 13, keep=set()),
        'total': row_spec(wd, 11, 13, keep=set()),
        'widths': col_widths(wd, 13),
        'orientation': dps.orientation or 'landscape',
        'zoom': 80,
    }

    spec = {'styles': styles, 'header': header, 'headerMerges': merges, 'roles': roles,
            'widths': widths, 'page': page, 'rate': rate, 'cost': cost, 'daily': daily}
    out = Path(__file__).resolve().parent.parent / 'template-spec.js'
    out.write_text('// 由 tools/extract_template_spec.py 自動產生，請勿手改\n'
                   '// 來源：公司範本「樣本-出差報告暨旅費核銷.xlsx」\n'
                   'window.TEMPLATE_SPEC = ' + json.dumps(spec, ensure_ascii=False, separators=(',', ':')) + ';\n',
                   encoding='utf-8')
    print(f'wrote {out} ({len(styles)} styles)')


if __name__ == '__main__':
    main(sys.argv[1] if len(sys.argv) > 1 else '樣本-出差報告暨旅費核銷.xlsx')
